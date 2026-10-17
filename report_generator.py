"""
report_generator.py
-------------------
Generates a Word (.docx) report from a comparison result dict, following the
table structure required by 「test_report_prompt」 section 7 审查结果输出格式:
  - Overall conclusion (符合 / 不符合 / 需人工复核后确认)
  - Issue list table (序号 / 问题描述 / 测试要求的内容 / 测试报告的内容 / 问题归类 / 严重程度)
  - Manual-review image list (待人工核查图片清单)
  - Recommendations (建议)
"""
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.section import WD_ORIENT


# ── Helpers ────────────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    """Set background colour of a table cell."""
    from docx.oxml import parse_xml
    shading = parse_xml(
        f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'w:fill="{hex_color}" w:val="clear"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading)


def _set_landscape(doc: Document):
    """Switch the current section to landscape (the wide issue table needs more width)."""
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width


def _header_row(table, headers: list):
    """Fill the first row of a table as a dark header row with white bold text."""
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        _set_cell_bg(cell, "595959")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _shade_alternate_row(cells, idx: int, skip: set = frozenset()):
    """Light-gray shading for even data rows, for readability."""
    if idx % 2 == 0:
        for i, cell in enumerate(cells):
            if i not in skip:
                _set_cell_bg(cell, "F2F2F2")


OVERALL_COLORS = {
    "符合":            "C6EFCE",  # green
    "不符合":          "FFC7CE",  # red
    "需人工复核后确认":  "FFEB9C",  # yellow
    "error":          "D9D9D9",  # gray – system/API failure, not a real verdict
}

SEVERITY_COLORS = {
    "严重":     "FFC7CE",  # red
    "一般":     "FCE4D6",  # orange
    "轻微":     "FFEB9C",  # yellow
    "需人工复核": "DDEBF7",  # blue
}


def _overall_color(overall: str) -> str:
    return OVERALL_COLORS.get(overall, OVERALL_COLORS.get(overall.lower(), "FFFFFF"))


def _severity_color(severity: str) -> str:
    return SEVERITY_COLORS.get(severity.strip(), "FFFFFF")


# ── Main Generator ─────────────────────────────────────────────────────────────

def generate_word_report(result: dict, tc_name: str, tr_name: str) -> bytes:
    """
    Generate a structured Word document from a comparison result.
    Layout:
      1. Title
      2. Info table              (metadata, 2-col)
      3. Summary                 (one-line paragraph)
      4. Issue list table        (序号 / 问题描述 / 测试要求的内容 / 测试报告的内容 / 问题归类 / 严重程度)
      5. Manual-review image list (待人工核查图片清单, if any)
      6. Recommendations          (建议)
    """
    doc = Document()
    _set_landscape(doc)

    # ── 1. Title ───────────────────────────────────────────────────
    doc.add_heading("Test Report Review Result / 测试报告审查结果", level=0)

    # ── 2. Info table ──────────────────────────────────────────────
    overall = result.get("overall", "N/A")
    meta_rows = [
        ("Test Case",         tc_name),
        ("Test Report",       tr_name),
        ("Generated",         datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Overall / 总体结论", overall),
    ]
    info_tbl = doc.add_table(rows=len(meta_rows), cols=2)
    info_tbl.style = "Table Grid"
    for i, (label, value) in enumerate(meta_rows):
        lc, vc = info_tbl.rows[i].cells
        lc.text = label
        lc.paragraphs[0].runs[0].bold = True
        _set_cell_bg(lc, "D9D9D9")
        vc.text = value
        if label.startswith("Overall"):
            _set_cell_bg(vc, _overall_color(overall))
            vc.paragraphs[0].runs[0].bold = True
    for row in info_tbl.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(7.5)

    doc.add_paragraph()  # spacer

    # ── 3. Summary ─────────────────────────────────────────────────
    doc.add_heading("Summary / 总结", level=1)
    doc.add_paragraph(result.get("summary", "—"))

    # ── 4. Issue list table ──────────────────────────────────────────
    doc.add_heading("Issue List / 问题清单", level=1)
    issues = result.get("issues", [])

    if issues:
        tbl = doc.add_table(rows=1, cols=6)
        tbl.style = "Table Grid"
        _header_row(tbl, [
            "#", "问题描述 / Description", "测试要求的内容 / Requirement",
            "测试报告的内容 / Report Content", "问题归类 / Category", "严重程度 / Severity",
        ])

        for idx, issue in enumerate(issues, 1):
            row_cells = tbl.add_row().cells
            row_cells[0].text = str(issue.get("no", idx))
            row_cells[1].text = issue.get("description", "")
            row_cells[2].text = issue.get("requirement_text", "")
            row_cells[3].text = issue.get("report_text", "")
            row_cells[4].text = issue.get("category", "")
            severity = issue.get("severity", "")
            row_cells[5].text = severity
            if row_cells[5].paragraphs[0].runs:
                row_cells[5].paragraphs[0].runs[0].bold = True
            _set_cell_bg(row_cells[5], _severity_color(severity))
            _shade_alternate_row(row_cells, idx, skip={5})

        col_widths = [Inches(0.35), Inches(2.3), Inches(2.1), Inches(2.1), Inches(1.5), Inches(1.15)]
        for row in tbl.rows:
            for cell, w in zip(row.cells, col_widths):
                cell.width = w
    else:
        doc.add_paragraph("(No issues found / 未发现问题)")
        details = result.get("details", "").strip()
        if details:
            doc.add_heading("Analysis Notes / 分析说明", level=2)
            doc.add_paragraph(details)

    # ── 5. Manual-review image list ───────────────────────────────
    manual_images = result.get("manual_review_images", [])
    if manual_images:
        doc.add_heading("Manual Review Image List / 待人工核查图片清单", level=1)
        mtbl = doc.add_table(rows=1, cols=4)
        mtbl.style = "Table Grid"
        _header_row(mtbl, [
            "位置 / Location", "对应要求 / Requirement",
            "报告描述 / Report Description", "无法识别原因 / Reason",
        ])

        for idx, img in enumerate(manual_images, 1):
            rc = mtbl.add_row().cells
            rc[0].text = img.get("location", "")
            rc[1].text = img.get("requirement_text", "")
            rc[2].text = img.get("report_text", "")
            rc[3].text = img.get("reason", "")
            _shade_alternate_row(rc, idx)

        col_widths = [Inches(1.8), Inches(2.5), Inches(2.5), Inches(2.2)]
        for row in mtbl.rows:
            for cell, w in zip(row.cells, col_widths):
                cell.width = w

    # ── 6. Recommendations ────────────────────────────────────────
    recommendations = result.get("recommendations", "").strip()
    if recommendations:
        doc.add_heading("Recommendations / 建议", level=1)
        doc.add_paragraph(recommendations)

    # ── Serialize ──────────────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()